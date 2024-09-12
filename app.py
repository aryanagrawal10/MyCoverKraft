import streamlit as st
from datetime import date
from io import BytesIO
from fpdf import FPDF
from docx import Document

# Placeholder function definitions
def extract_text_from_pdf(pdf_file):
    return "Extracted text from the uploaded PDF file."

def ai_chat_completion(*args, **kwargs):
    return {
        'choices': [
            {'message': {'content': "Generated cover letter based on provided inputs."}}
        ]
    }

def save_feedback_to_file(user_name, feedback):
    pass

def match_keywords(resume_text, job_description):
    return "Matched Keywords"

# Streamlit UI setup
st.markdown("# 📝 MyCoverKraft - Your Personalized Cover Letter Generator")

tab1, tab2, tab3 = st.tabs(["Cover Letter Generator", "Resume Parser and Editor", "Resume and Job Description Keyword Matcher"])

if 'cover_letter_generated' not in st.session_state:
    st.session_state.cover_letter_generated = False
if 'feedback_submitted' not in st.session_state:
    st.session_state.feedback_submitted = False

# Cover Letter Generator Tab
with tab1:
    st.markdown("## Cover Letter Generator")
    st.expander("Instructions").write("""
        - Fields marked with an asterisk (*) are mandatory.
        - Upload your resume or copy your resume/experiences.
        - Paste a relevant job description.
        - Input other relevant data.
    """)

    res_format = st.radio(
        "Resume Input Method",
        ('Upload', 'Paste'),
        help="Choose how you'd like to input your resume."
    )

    res_text = ""
    if res_format == 'Upload':
        res_file = st.file_uploader('📁 Upload your resume in PDF format', type='pdf', key="res_file_upload")
        if res_file:
            res_text = extract_text_from_pdf(res_file)
    else:
        res_text = st.text_area('Pasted resume elements', key="pasted_resume_elements")

    st.info("Your data privacy is important. Uploaded resumes are only used for generating the cover letter and are not stored or used for any other purposes.")

    tone = st.selectbox('Select the Tone of Your Cover Letter', ['Professional', 'Friendly', 'Enthusiastic', 'Formal', 'Casual'], key="tone_selectbox")
    achievements = st.text_area('Include Specific Achievements, Skills or Keywords', key="achievements_textarea")
    letter_structure = st.radio('Choose Your Cover Letter Structure', ('Standard', 'Skill-based', 'Story-telling'), key="letter_structure_radio")

    with st.form('input_form'):
        job_desc = st.text_area('Job description*', key="job_desc_textarea")
        user_name = st.text_input('Name*', key="user_name_input")
        company = st.text_input('Company name*', key="company_name_input")
        manager = st.text_input('Hiring manager', key="manager_input")
        role = st.text_input('Job Role*', key="role_input")
        referral = st.text_input('How did you find out about this opportunity?', key="referral_input")

        submitted = st.form_submit_button("Generate Cover Letter")

    if submitted:
        if res_text and job_desc and user_name and company and role:
            try:
                customization_prompt = f"""
                Tone: {tone.lower()}
                Achievements/Skills: {achievements}
                Structure: {letter_structure.lower()}
                """
                completion = ai_chat_completion(
                    model="gpt-3.5-turbo",
                    temperature=0.99,
                    messages=[
                        {"role": "user", "content": "You will need to generate a cover letter based on specific resume and a job description"},
                        {"role": "user", "content": f"My resume text: {res_text}"},
                        {"role": "user", "content": f"The job description is: {job_desc}"},
                        {"role": "user", "content": f"The candidate's name to include on the cover letter: {user_name}"},
                        {"role": "user", "content": f"The job title/role : {role}"},
                        {"role": "user", "content": f"The hiring manager is: {manager}"},
                        {"role": "user", "content": f"How you heard about the opportunity: {referral}"},
                        {"role": "user", "content": f"The company to which you are generating the cover letter for: {company}"},
                        {"role": "user", "content": customization_prompt},
                        {"role": "user", "content": "The cover letter should have three content paragraphs. Please replace all placeholders with the specific details provided."},
                        {"role": "user", "content": "Generate a specific cover letter based on the above. Generate the response and include appropriate spacing between the paragraph text."}
                    ]
                )

                response_out = completion['choices'][0]['message']['content']
                response_out = response_out.replace('[Hiring Manager]', manager if manager else 'Hiring Manager')
                response_out = response_out.replace('[Recipient\'s Name]', manager if manager else 'Hiring Manager')
                response_out = response_out.replace('[Your Name]', user_name if user_name else 'Your Name')
                response_out = response_out.replace('[Job description*]', job_desc)
                response_out = response_out.replace('[Company Name]', company)
                response_out = response_out.replace('[Job Role*]', role)
                today_date = date.today().strftime("%B %d, %Y")
                response_out = response_out.replace('[Today\'s Date]', today_date)
                response_out = response_out.replace('[Today’s Date]', today_date)
                response_out = response_out.replace('[Date]', today_date)
                response_out = response_out.replace('[Company Address]', '')
                response_out = response_out.replace('[Your Address]', '')
                response_out = response_out.replace('[City, State, ZIP Code]', '')
                response_out = response_out.replace('[City, State, ZIP]', '')
                response_out = response_out.replace('[Email Address]', '')
                response_out = response_out.replace('[Phone Number]', '')
                response_out = response_out.replace('[Your Contact Information]', '')
                response_out = response_out.replace('[How did you find out about this opportunity?]', referral)

                st.write(response_out)
                
                # Save the generated cover letter in different formats
                def create_file(data, file_type):
                    if file_type == 'txt':
                        return data.encode('utf-8')
                    elif file_type == 'docx':
                        doc = Document()
                        doc.add_paragraph(data)
                        buffer = BytesIO()
                        doc.save(buffer)
                        return buffer.getvalue()
                    elif file_type == 'pdf':
                        pdf = FPDF()
                        pdf.add_page()
                        pdf.set_font("Arial", size=12)
                        pdf.multi_cell(0, 10, data)
                        buffer = BytesIO()
                        pdf.output(dest='S').encode('latin-1')
                        buffer.write(pdf.output(dest='S').encode('latin-1'))
                        return buffer.getvalue()
                
                col1, col2, col3 = st.columns(3)

                with col1:
                    st.download_button('Download TXT', create_file(response_out, 'txt'), f'{user_name}_cover_letter.txt', 'text/plain', key="download_txt")
                with col2:
                    st.download_button('Download DOCX', create_file(response_out, 'docx'), f'{user_name}_cover_letter.docx', 'application/vnd.openxmlformats-officedocument.wordprocessingml.document', key="download_docx")
                with col3:
                    st.download_button('Download PDF', create_file(response_out, 'pdf'), f'{user_name}_cover_letter.pdf', 'application/pdf', key="download_pdf")

                st.session_state.cover_letter_generated = True
            except Exception as e:
                st.error(f"An error occurred while generating the cover letter: {e}")
        else:
            st.error("Please fill in all the required fields.")

    if st.session_state.cover_letter_generated and not st.session_state.feedback_submitted:
        feedback = st.slider("Rate the quality of the generated cover letter (1-5)", 1, 5, 3, key="feedback_slider")
        if st.button("Submit Feedback", key="feedback_button"):
            save_feedback_to_file(user_name, feedback)
            st.success("Thank you for your feedback!")
            st.session_state.feedback_submitted = True

# Resume Parser and Editor Tab
with tab2:
    st.markdown("## Resume Parser and Editor")
    st.expander("Instructions").write("""
        - Upload your resume in PDF format.
        - Extracted text will appear in the text area.
        - Edit the extracted text as needed.
    """)

    resume_file = st.file_uploader('📁 Upload your resume in PDF format', type='pdf', key="resume_file_editor")
    if resume_file:
        resume_text = extract_text_from_pdf(resume_file)
        st.text_area('Extracted Resume Text', value=resume_text, height=300, key="resume_text_area")

        edited_text = st.text_area('Edit Resume Text', value=resume_text, height=300, key="edit_resume_text_area")
        if st.button("Save Edited Resume", key="save_edited_resume"):
            st.session_state.edited_resume_text = edited_text
            st.success("Edited resume text saved successfully!")

# Resume and Job Description Keyword Matcher Tab
with tab3:
    st.markdown("## Resume and Job Description Keyword Matcher")
    st.expander("Instructions").write("""
        - Upload your resume or paste it.
        - Paste a job description.
        - The app will highlight matching keywords between your resume and the job description.
    """)

    resume_input_method = st.radio(
        "Resume Input Method",
        ('Upload', 'Paste'),
        help="Choose how you'd like to input your resume.",
        key="resume_input_method_matcher"
    )

    resume_text_matcher = ""
    if resume_input_method == 'Upload':
        resume_file_matcher = st.file_uploader('📁 Upload your resume in PDF format', type='pdf', key="resume_file_matcher")
        if resume_file_matcher:
            resume_text_matcher = extract_text_from_pdf(resume_file_matcher)
    else:
        resume_text_matcher = st.text_area('Pasted resume elements', key="resume_paste_text_area")

    job_description_matcher = st.text_area('Job description', key="job_description_matcher_textarea")
    if st.button("Match Keywords", key="match_keywords_button"):
        matched_keywords = match_keywords(resume_text_matcher, job_description_matcher)
        st.write(f"Matched Keywords: {matched_keywords}")
